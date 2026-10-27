#!/usr/bin/env python3
"""Build docs/examples/example-readout.pptx, the example readout deck.

Usage: python3 tools/examples/build_readout_deck.py docs/examples/example-readout.pptx

Needs python-pptx (pip install python-pptx). The text matches
docs/examples/example-readout.md; change both together.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

out = sys.argv[1]
INK = RGBColor(0x1F, 0x2A, 0x24); GREEN = RGBColor(0x2C, 0x5F, 0x2D); LEAF = RGBColor(0x97, 0xBC, 0x62)
MIST = RGBColor(0xEE, 0xF3, 0xEC); SLATE = RGBColor(0x50, 0x80, 0x8E); MUTED = RGBColor(0x5B, 0x6A, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); AMBER = RGBColor(0x9A, 0x5B, 0x00); AMBERBG = RGBColor(0xFB, 0xEF, 0xD9)
LINE = RGBColor(0xD5, 0xDE, 0xD3)
SERIF, SANS = "Cambria", "Calibri"
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.7)

prs = Presentation(); prs.slide_width = W; prs.slide_height = H
BLANK = prs.slide_layouts[6]
TITLES = []

def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    return sh

def text(s, x, y, w, h, runs, size=18, color=INK, font=SANS, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1, italic=False):
    """runs: str or list of paragraphs; a paragraph is str or list of (text, overrides)."""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        if i: p.space_before = Pt(size * 0.5)
        for chunk in (para if isinstance(para, list) else [(para, {})]):
            t, o = chunk if isinstance(chunk, tuple) else (chunk, {})
            r = p.add_run(); r.text = t; f = r.font
            f.name = o.get("font", font); f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic); f.color.rgb = o.get("color", color)
    return tb

def slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = GREEN if dark else WHITE
    return s

def header(s, label, title, dark=False, n=None):
    TITLES.append(title)
    text(s, M, Inches(0.55), Inches(11.5), Inches(0.3), label.upper(), size=12, bold=True, color=LEAF if dark else SLATE)
    text(s, M, Inches(0.85), Inches(11.9), Inches(0.7), title, size=32, font=SERIF, bold=True, color=WHITE if dark else GREEN)
    foot = "Fictional example · made with Magpie"
    text(s, M, Inches(7.0), Inches(6), Inches(0.25), foot, size=10, color=LEAF if dark else MUTED)
    if n: text(s, Inches(12.0), Inches(7.0), Inches(0.6), Inches(0.25), str(n), size=10, color=LEAF if dark else MUTED, align=PP_ALIGN.RIGHT)

def chip(s, x, y, label, fill=MIST, color=GREEN, w=None):
    w = w or Inches(0.13 * len(label) + 0.45)
    b = box(s, x, y, w, Inches(0.38), fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = b.text_frame; tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = label
    r.font.size = Pt(13); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = color
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return w

def notes(s, t): s.notes_slide.notes_text_frame.text = t

n = 0
# 1 Title
s = slide(dark=True); n += 1; TITLES.append("Export isn't the problem. Import is.")
text(s, M, Inches(2.3), Inches(11), Inches(0.35), "UX RESEARCH READOUT · LEDGERLINE", size=13, bold=True, color=LEAF)
text(s, M, Inches(2.75), Inches(12), Inches(1.0), "Export isn't the problem. Import is.", size=48, font=SERIF, bold=True, color=WHITE)
text(s, M, Inches(3.95), Inches(11), Inches(1.2), ["Why new customers give up on export, and what to do about it",
     [("14 October 2026 · Priya Nair, UX Research · For the Reporting team", {"size": 16, "color": LEAF})]], size=22, color=WHITE)
text(s, M, Inches(7.0), Inches(6), Inches(0.25), "Fictional example · made with Magpie", size=10, color=LEAF)
notes(s, "Open with the answer. The rest of the deck shows how we know.")

# 2 The short version: drawn from this study's focus and findings, not fixed headings
s = slide(); n += 1; header(s, "The short version", "What we learned, in short", n=n)
text(s, M, Inches(2.0), Inches(4.6), Inches(3.6), [[("We set out to learn why new customers give up on export.", {"font": SERIF, "size": 26, "bold": True, "color": GREEN})],
     [("The answer points earlier in the journey than we expected.", {"size": 18, "color": MUTED})]], size=26, spacing=1.15)
points = [("Import is where it breaks.", "For 5 of the 6 people who gave up, import had quietly skipped columns. A new export screen wouldn't have helped them."),
          ("The problem is bigger than our numbers show.", "6 of 8 fix the file by emailing it to a colleague, so these exports count as successes."),
          ("Finance teams are fully blocked.", "Only 2 people, but neither could get a file their finance system accepts.")]
for i, (h, b) in enumerate(points):
    y = Inches(1.95) + i * Inches(1.5)
    box(s, Inches(5.75), y, Inches(6.9), Inches(1.32), fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, Inches(6.1), y + Inches(0.2), Inches(6.3), Inches(0.4), h, size=19, bold=True, color=GREEN)
    text(s, Inches(6.1), y + Inches(0.6), Inches(6.3), Inches(0.7), b, size=15, spacing=1.1)
text(s, M, Inches(6.45), Inches(11.9), Inches(0.4), [[("So: ", {"bold": True, "color": GREEN}), ("fix import first, then decide whether the export screen still needs a rebuild.", {})]], size=18)
notes(s, "If people only remember one slide, it's this one.")

# 3 The decision
s = slide(); n += 1; header(s, "Why we did this", "The team is choosing where Q1 effort goes", n=n)
for i, (h, b) in enumerate([("Option A", "Rebuild the export screen"), ("Option B", "Fix the import step that feeds it")]):
    x = M + i * Inches(4.15)
    box(s, x, Inches(2.0), Inches(3.9), Inches(1.9), fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, x + Inches(0.35), Inches(2.3), Inches(3.3), Inches(0.35), h.upper(), size=12, bold=True, color=SLATE)
    text(s, x + Inches(0.35), Inches(2.7), Inches(3.3), Inches(1.0), b, size=24, font=SERIF, bold=True, color=GREEN)
text(s, M, Inches(4.35), Inches(8.0), Inches(2.4), [
    [("Decision by: ", {"bold": True}), ("30 October, by the Reporting team.", {})],
    [("Why it matters: ", {"bold": True}), ("212 export tickets last quarter, and more weekly sharing is the team's goal for the second half of the year.", {})],
    [("What we asked: ", {"bold": True}), ("Why do new customers give up on export in their first 90 days?", {})]], size=17, spacing=1.15)
box(s, Inches(9.25), Inches(2.0), Inches(3.4), Inches(2.9), fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, Inches(9.6), Inches(2.3), Inches(2.8), Inches(3.9), [[("What we already knew", {"font": SERIF, "bold": True, "size": 18, "color": GREEN})],
     "Export is the second most-used admin feature (2025 survey).", "Field reps hear that \"customers hate the export\".", "Nobody had asked why it fails."], size=15, spacing=1.15)

# 4 What we did
s = slide(); n += 1; header(s, "What we did", "We watched 8 customers try a real export", n=n)
tiles = [("8", "customers, one at a time"), ("45 min", "each, by video call"), ("< 90 days", "since they signed up"), ("6–10 Oct", "2026")]
tw = Inches(2.8)
for i, (big, small) in enumerate(tiles):
    x = M + i * (tw + Inches(0.2))
    box(s, x, Inches(2.0), tw, Inches(1.9), fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, x, Inches(2.3), tw, Inches(0.8), big, size=40, font=SERIF, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.15), tw, Inches(0.5), small, size=15, color=MUTED, align=PP_ALIGN.CENTER)
text(s, M, Inches(4.4), Inches(11.8), Inches(2.0), [
    "Each person told us about the last export they tried, then exported last month's revenue report on their own account while we watched.",
    "They all build and share reports at companies with 50 to 500 Ledgerline seats, and our Customer Success team helped us find them."], size=18, spacing=1.2)

# 5 Hook
s = slide(dark=True); n += 1; header(s, "The story", "We expected the export screen to be the problem", dark=True)
text(s, M, Inches(2.6), Inches(11), Inches(2.5), ["The study was set up to help plan a new export screen. Field reps said customers hate it.",
     [("Watching people use it told a different story.", {"color": LEAF, "bold": True})]], size=26, color=WHITE, spacing=1.2)

# 6 Journey
s = slide(); n += 1; header(s, "Where it goes wrong", "The trouble starts before export", n=n)
steps = ["Import data", "Build a report", "Export the file", "Send to finance"]
sw = Inches(2.65); sy = Inches(2.2); sx = [M + i * (sw + Inches(0.42)) for i in range(4)]
for i, st in enumerate(steps):
    b = box(s, sx[i], sy, sw, Inches(0.95), fill=GREEN if i in (0,) else MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    text(s, sx[i], sy + Inches(0.08), sw, Inches(0.3), f"STEP {i+1}", size=11, bold=True, color=LEAF if i == 0 else SLATE, align=PP_ALIGN.CENTER)
    text(s, sx[i], sy + Inches(0.38), sw, Inches(0.45), st, size=20, font=SERIF, bold=True, color=WHITE if i == 0 else GREEN, align=PP_ALIGN.CENTER)
    if i < 3:
        a = box(s, sx[i] + sw + Inches(0.08), sy + Inches(0.3), Inches(0.26), Inches(0.35), fill=LEAF, shape=MSO_SHAPE.CHEVRON)
callouts = [(0, "5 of 8", "Import skipped columns. Nobody noticed until the export came out with gaps."),
            (2, "7 of 8", "Didn't understand the \"Include hidden fields\" checkbox."),
            (3, "2 of 8", "Finance teams: their system rejected the file's column order."),
            (3, "6 of 8", "Fixed the file by emailing it to a colleague.")]
ys = {0: Inches(3.55), 2: Inches(3.55), 3: Inches(3.55)}
drawn = set()
for col, big, msg in callouts:
    y = ys[col]
    if col not in drawn:
        drawn.add(col)
        box(s, sx[col] + sw / 2 - Emu(6350), sy + Inches(0.95), Emu(12700), y - sy - Inches(0.95), fill=LINE)
    text(s, sx[col], y, sw, Inches(0.5), big, size=28, font=SERIF, bold=True, color=GREEN)
    text(s, sx[col], y + Inches(0.55), sw, Inches(1.2), msg, size=15, spacing=1.15)
    ys[col] = y + Inches(1.65)
text(s, sx[1], Inches(3.55), sw, Inches(1.6), "No problems here.", size=15, color=MUTED, italic=True)

# 7-10 findings
FIND = [
 ("Finding 1 of 4", "5 of 8 hit columns that import had skipped", "5 of 8", "Strong evidence",
  "When the export preview showed empty columns, they stopped. In every case, import had skipped those columns weeks earlier.",
  "I thought export was broken, but it's the import that never had them.", "P3, operations analyst, 120 seats",
  "P2 and P5 had no gaps. Customer Success had set up their import for them.",
  "A new export screen wouldn't have helped 5 of the 6 people who gave up.", "P1, P3, P4, P6, P8", None),
 ("Finding 2 of 4", "Finance teams can't use the file at all", "2 of 8", "Weak evidence",
  "Both people from finance teams need columns in a fixed order for their finance system. Export can't do that, so the file is rejected.",
  "If the order's wrong the reconciliation tool rejects the whole file.", "P7, finance controller, 450 seats",
  "The other 6 didn't need a fixed order. This may only affect finance teams.",
  "A small group is completely stuck. Support can count how many customers this affects.", "P5, P7", "Serious: fully blocked"),
 ("Finding 3 of 4", "6 of 8 email the file to someone to fix", "6 of 8", "Strong evidence",
  "When the file is wrong, people send it to a colleague to fix by hand. They told us this without being asked.",
  "Sent the raw file to [name] and asked them to fix it.", "P1, operations manager, 80 seats",
  "P2 and P5 didn't do this. P2 rarely exports.",
  "Our numbers count these exports as successes, so the 212 tickets undercount the problem.", "P1, P3, P4, P6, P7, P8", None),
 ("Finding 4 of 4", "7 of 8 didn't understand one checkbox", "7 of 8", "Moderate evidence",
  "Almost everyone skipped \"Include hidden fields\" or guessed what it does. Nobody could say what a hidden field is.",
  "Hidden from who?", "P4, reporting lead, 310 seats",
  "P8 used it correctly after reading the help article.",
  "A quick wording fix on the export screen. It isn't why people give up.", "P1 to P7", None),
]
for lab, title, stat, ev, what, quote, who, cuts, means, heard, flag in FIND:
    s = slide(); n += 1; header(s, lab, title, n=n)
    box(s, M, Inches(1.95), Inches(3.6), Inches(4.7), fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, M, Inches(2.35), Inches(3.6), Inches(1.2), stat, size=66, font=SERIF, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    text(s, M + Inches(0.3), Inches(3.65), Inches(3.0), Inches(0.4), "people we watched", size=15, color=MUTED, align=PP_ALIGN.CENTER)
    cw2 = chip(s, M + Inches(0.5), Inches(4.4), ev, w=Inches(2.6))
    if flag: chip(s, M + Inches(0.5), Inches(4.95), flag, fill=AMBERBG, color=AMBER, w=Inches(2.6))
    text(s, M + Inches(0.3), Inches(5.75), Inches(3.0), Inches(0.7), [[("Heard from: ", {"bold": True}), (heard, {})]], size=13, color=MUTED, align=PP_ALIGN.CENTER)
    X = Inches(4.75); RW = Inches(7.9)
    text(s, X, Inches(1.95), RW, Inches(1.0), what, size=19, spacing=1.2)
    box(s, X, Inches(3.1), Emu(50800), Inches(1.0), fill=LEAF)
    text(s, X + Inches(0.3), Inches(3.1), RW - Inches(0.3), Inches(1.0), [[(f"“{quote}”", {"font": SERIF, "size": 21, "color": GREEN})],
         [(who, {"size": 13, "color": MUTED})]], size=21)
    for j, (h, b) in enumerate([("Who didn't fit", cuts), ("What it means", means)]):
        x = X + j * Inches(4.05)
        text(s, x, Inches(4.6), Inches(3.8), Inches(0.3), h.upper(), size=12, bold=True, color=SLATE)
        text(s, x, Inches(4.95), Inches(3.8), Inches(1.6), b, size=16, spacing=1.15)

# 11 Other findings
s = slide(); n += 1; header(s, "Other findings", "Three smaller things we noticed", n=n)
others = [("2 of 8", "expected export to run every month by itself.", "Not something we asked about. Worth raising in roadmap planning."),
          ("3 of 8", "picked the same columns every time.", "The export screen doesn't remember last time's choice."),
          ("3 of 8", "renamed the file before sending it.", "Nobody gave up over this, so it isn't a priority.")]
for i, (big, msg, sub) in enumerate(others):
    x = M + i * (Inches(3.8) + Inches(0.25))
    box(s, x, Inches(2.0), Inches(3.8), Inches(3.6), fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, x + Inches(0.35), Inches(2.3), Inches(3.1), Inches(0.7), big, size=34, font=SERIF, bold=True, color=GREEN)
    text(s, x + Inches(0.35), Inches(3.1), Inches(3.1), Inches(1.0), msg, size=18, bold=True)
    text(s, x + Inches(0.35), Inches(4.15), Inches(3.1), Inches(1.2), sub, size=15, color=MUTED, spacing=1.15)
text(s, M, Inches(6.0), Inches(11.8), Inches(0.5), "All weak evidence. Every finding from the study is in this deck or the report's appendix.", size=14, color=MUTED)

# 12 Recommendations
s = slide(); n += 1; header(s, "What to do", "Four things to do, in this order", n=n)
recs = [("1", "Warn people at import when columns their reports use were skipped.", "Data Pipeline, Reporting", "Fixes finding 1"),
        ("2", "Hold the export screen rebuild until the import warning is live and measured.", "Reporting", "Based on findings 1 and 3"),
        ("3", "Count how many finance-team customers need a fixed column order.", "Support, Customer Success", "Sizes finding 2"),
        ("4", "Rename \"Include hidden fields\" and add a one-line explanation.", "Design", "Fixes finding 4")]
for i, (num, act, owner, why) in enumerate(recs):
    x = M + (i % 2) * Inches(6.05); y = Inches(1.95) + (i // 2) * Inches(2.35)
    box(s, x, y, Inches(5.85), Inches(2.15), fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    c = box(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55), Inches(0.55), fill=GREEN, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.margin_left = tf.margin_right = 0; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = SANS
    text(s, x + Inches(1.1), y + Inches(0.3), Inches(4.5), Inches(1.0), act, size=18, bold=True, spacing=1.15)
    text(s, x + Inches(1.1), y + Inches(1.35), Inches(4.5), Inches(0.7), [[("Teams involved: ", {"bold": True}), (owner, {})], [(why, {"color": MUTED})]], size=14)

text(s, M, Inches(6.55), Inches(11.9), Inches(0.35), "Who takes each one is for the teams to agree together.", size=14, color=MUTED)

# 13 Next
s = slide(); n += 1; header(s, "What happens next", "The team decides by 30 October", n=n)
steps = [("24 Oct", "Support counts the finance-team customers", "So we know how big finding 2 is."),
         ("30 Oct", "The team decides", "Using this deck and the report."),
         ("November", "3 follow-up sessions", "With finance teams, if the import warning goes ahead."),
         ("Later", "One open question", "Does the import problem affect scheduled reports too?")]
lw = Inches(2.85)
box(s, M + Inches(0.2), Inches(2.42), Inches(11.6), Emu(19050), fill=LINE)
for i, (when, what, sub) in enumerate(steps):
    x = M + i * (lw + Inches(0.17))
    box(s, x, Inches(2.25), Inches(0.38), Inches(0.38), fill=LEAF if i else GREEN, shape=MSO_SHAPE.OVAL)
    text(s, x, Inches(2.9), lw, Inches(0.5), when, size=24, font=SERIF, bold=True, color=GREEN)
    text(s, x, Inches(3.5), lw - Inches(0.2), Inches(0.9), what, size=17, bold=True, spacing=1.1)
    text(s, x, Inches(4.45), lw - Inches(0.2), Inches(1.2), sub, size=15, color=MUTED, spacing=1.15)

# 14 Limits
s = slide(); n += 1; header(s, "Read with care", "What this study can't tell us", n=n)
text(s, M, Inches(2.0), Inches(7.6), Inches(4.2), [
    [("Who we spoke to. ", {"bold": True}), ("Customer Success helped find all 8 people. Customers who set up on their own may differ.", {})],
    [("Finance teams. ", {"bold": True}), ("Only 2 of the 8, so finding 2 could be bigger or smaller than it looks.", {})],
    [("How many. ", {"bold": True}), ("We didn't check the export numbers first, so we can't say what share of all customers hit finding 1.", {})]], size=18, spacing=1.2)
box(s, Inches(8.9), Inches(2.0), Inches(3.75), Inches(2.4), fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, Inches(9.2), Inches(2.3), Inches(3.15), Inches(1.9), "Magpie drafted this slide and the researcher chose to keep it. The limits are in the report either way.", size=14, color=MUTED, italic=True, spacing=1.15)

# 15 Appendix: who
s = slide(); n += 1; header(s, "Appendix", "Who we spoke to", n=n)
rows = [("Code", "Role", "Seats", "Team"), ("P1", "Operations manager", "80", "Operations"), ("P2", "Office manager", "60", "Operations"),
        ("P3", "Operations analyst", "120", "Operations"), ("P4", "Reporting lead", "310", "Reporting"), ("P5", "Finance analyst", "200", "Finance"),
        ("P6", "Account manager", "95", "Sales"), ("P7", "Finance controller", "450", "Finance"), ("P8", "Data analyst", "150", "Reporting")]
tbl = s.shapes.add_table(len(rows), 4, M, Inches(1.9), Inches(9.0), Inches(4.6)).table
for j, wd in enumerate([1.2, 3.6, 1.6, 2.6]): tbl.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = v
        c.fill.solid(); c.fill.fore_color.rgb = GREEN if i == 0 else (MIST if i % 2 else WHITE)
        r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(15); r.font.name = SANS
        r.font.bold = i == 0; r.font.color.rgb = WHITE if i == 0 else INK

# 16 Appendix: AI
s = slide(); n += 1; header(s, "Appendix", "How AI was used", n=n)
text(s, M, Inches(2.0), Inches(11.5), Inches(4.6), [
    [("Magpie, an AI research assistant, ", {"bold": True}), ("drafted the plan, the interview guide, the findings, the report and this deck.", {})],
    [("The researcher ", {"bold": True}), ("ran every session, noted their own patterns before seeing Magpie's, decided on each finding, and wrote the recommendations. They reviewed the findings on 12 October and this deck on 14 October.", {})],
    [("Checks: ", {"bold": True}), ("quotes and counts were checked by code against the transcripts. A separate AI reviewer raised 11 points, and 8 changed the findings. That reviewer can be wrong too.", {})],
    [("Transcripts ", {"bold": True}), ("came from the video tool, and the researcher corrected them against the recordings.", {})]], size=18, spacing=1.2)

prs.save(out)
long = [t for t in TITLES if len(t) > 50]
print(f"Saved {out}: {len(prs.slides)} slides")
if long:
    print("Titles that may not fit on one line:", long)
